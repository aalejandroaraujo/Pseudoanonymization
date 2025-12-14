"""
File Handler Module

Handles extraction of text from PDF, DOCX, and TXT files.
Preserves document structure (newlines, paragraphs) where possible.
"""

import os
import logging
from pathlib import Path
from typing import Optional

import pdfplumber
from docx import Document

logger = logging.getLogger(__name__)

SUPPORTED_FORMATS = {'.pdf', '.docx', '.txt'}


class FileHandlerError(Exception):
    """Custom exception for file handling errors."""
    pass


def validate_file(file_path: str) -> Path:
    """
    Validate that the file exists and has a supported format.

    Args:
        file_path: Path to the file to validate.

    Returns:
        Path object of the validated file.

    Raises:
        FileHandlerError: If file doesn't exist or format is unsupported.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileHandlerError(f"File not found: {file_path}")

    if not path.is_file():
        raise FileHandlerError(f"Path is not a file: {file_path}")

    if path.suffix.lower() not in SUPPORTED_FORMATS:
        raise FileHandlerError(
            f"Unsupported file format: {path.suffix}. "
            f"Supported formats: {', '.join(SUPPORTED_FORMATS)}"
        )

    return path


def extract_text_from_pdf(file_path: Path) -> str:
    """
    Extract text from a PDF file using pdfplumber.

    Args:
        file_path: Path to the PDF file.

    Returns:
        Extracted text as a single string with page separations.

    Raises:
        FileHandlerError: If PDF extraction fails.
    """
    logger.info(f"Extracting text from PDF: {file_path}")

    try:
        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            total_pages = len(pdf.pages)
            logger.info(f"PDF has {total_pages} pages")

            for i, page in enumerate(pdf.pages, 1):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
                    logger.debug(f"Extracted page {i}/{total_pages}")
                else:
                    logger.warning(f"Page {i} appears to be empty or image-only")

        return '\n\n'.join(text_parts)

    except Exception as e:
        raise FileHandlerError(f"Failed to extract text from PDF: {e}")


def extract_text_from_docx(file_path: Path) -> str:
    """
    Extract text from a DOCX file using python-docx.
    Extracts both paragraphs and table content.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        Extracted text as a single string.

    Raises:
        FileHandlerError: If DOCX extraction fails.
    """
    logger.info(f"Extracting text from DOCX: {file_path}")

    try:
        doc = Document(file_path)
        text_parts = []

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Extract tables
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                table_text.append(' | '.join(row_text))
            if table_text:
                text_parts.append('\n'.join(table_text))

        logger.info(f"Extracted {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables")
        return '\n\n'.join(text_parts)

    except Exception as e:
        raise FileHandlerError(f"Failed to extract text from DOCX: {e}")


def extract_text_from_txt(file_path: Path) -> str:
    """
    Extract text from a TXT file.

    Args:
        file_path: Path to the TXT file.

    Returns:
        File content as a string.

    Raises:
        FileHandlerError: If TXT reading fails.
    """
    logger.info(f"Reading text from TXT: {file_path}")

    try:
        # Try different encodings
        encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']

        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                logger.debug(f"Successfully read file with encoding: {encoding}")
                return content
            except UnicodeDecodeError:
                continue

        raise FileHandlerError(f"Could not decode file with any supported encoding")

    except FileHandlerError:
        raise
    except Exception as e:
        raise FileHandlerError(f"Failed to read TXT file: {e}")


def extract_text_from_file(file_path: str) -> str:
    """
    Extract text from PDF, DOCX, or TXT files.

    This is the main entry point for text extraction. It validates the file,
    determines its type, and delegates to the appropriate extraction function.

    Args:
        file_path: Path to the file to extract text from.

    Returns:
        Extracted text as a single string.

    Raises:
        FileHandlerError: If file validation or extraction fails.
    """
    path = validate_file(file_path)
    suffix = path.suffix.lower()

    if suffix == '.pdf':
        return extract_text_from_pdf(path)
    elif suffix == '.docx':
        return extract_text_from_docx(path)
    elif suffix == '.txt':
        return extract_text_from_txt(path)
    else:
        raise FileHandlerError(f"Unsupported file format: {suffix}")


def write_output_file(content: str, output_path: str, format: str = 'txt') -> None:
    """
    Write anonymized content to an output file.

    Args:
        content: The anonymized text content.
        output_path: Path to the output file.
        format: Output format ('txt' or 'md').

    Raises:
        FileHandlerError: If writing fails.
    """
    try:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        with open(path, 'w', encoding='utf-8') as f:
            f.write(content)

        logger.info(f"Output written to: {output_path}")

    except Exception as e:
        raise FileHandlerError(f"Failed to write output file: {e}")


def get_default_output_path(input_path: str, suffix: str = '_anon') -> str:
    """
    Generate a default output path based on the input path.

    Args:
        input_path: Path to the input file.
        suffix: Suffix to add before the extension.

    Returns:
        Default output path as string.
    """
    path = Path(input_path)
    return str(path.parent / f"{path.stem}{suffix}.txt")
